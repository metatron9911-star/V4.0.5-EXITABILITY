"""Offline unit/integration tests; no keys, orders or network calls."""
import copy
import io
import json
import tempfile
import unittest
from pathlib import Path
from engine import CFG, USDC_MINT, Store, Experiment, export_excel, now_ms

class Quotes:
    def __init__(self): self.multiple=1; self.available=True; self.half_override=None
    def quote(self,a,b,raw):
        if not self.available:return dict(ok=False,reason='OFFLINE')
        value=int(raw) if a==USDC_MINT else int(int(raw)*self.multiple)
        if a!=USDC_MINT and self.half_override is not None and int(raw)==2175000:
            value=self.half_override
        return dict(ok=True,ts_ms=now_ms(),min_raw=value,out_raw=value,impact_pct=0)

class TestV6(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory(); self.path=Path(self.tmp.name)/'test.db'
        self.db=Store(self.path); self.md=Quotes(); self.exp=Experiment(self.db,self.md)
    def tearDown(self):self.db.cx.close(); self.tmp.cleanup()
    def opened(self):
        self.assertEqual(self.exp.open_cohort('MINT','=TOKEN',1,6),'COHORT_OPENED')
        return self.db.snapshot()['positions'][0]
    def test_total_loss_budget_and_reserve(self):
        p=self.opened()
        self.assertAlmostEqual(p['cost']+p['reserve'],4.5)
        self.md.multiple=.001
        self.exp.exit_one(p['id'])
        p=self.db.snapshot()['positions'][0]
        self.assertEqual(p['remaining_raw'],'0')
        self.assertLessEqual(p['cost']-p['received'],4.5)
        self.assertEqual(self.db.snapshot()['positions'][1]['remaining_raw'],'4350000')
    def test_half_at_two_x_not_enough(self):
        p=self.opened(); self.md.multiple=2
        self.exp.exit_one(p['id'])
        self.assertFalse(self.db.snapshot()['positions'][0]['half_done'])
        self.md.multiple=2.1; self.exp.exit_one(p['id'])
        p=self.db.snapshot()['positions'][0]
        self.assertTrue(p['half_done']); self.assertGreaterEqual(p['received'],p['cost'])
        self.assertEqual(p['remaining_raw'],'2175000')
    def test_half_quote_not_full_quote_divided_by_two(self):
        p=self.opened(); self.md.multiple=2.2; self.md.half_override=4300000
        self.exp.exit_one(p['id'])
        self.assertFalse(self.db.snapshot()['positions'][0]['half_done'])
    def test_no_duplicate_half_after_restart(self):
        p=self.opened(); self.md.multiple=2.1; self.exp.exit_one(p['id'])
        self.db.cx.close(); self.db=Store(self.path); self.exp=Experiment(self.db,self.md)
        self.exp.exit_one(p['id'])
        p=self.db.snapshot()['positions'][0]
        self.assertEqual(p['remaining_raw'],'2175000'); self.assertAlmostEqual(p['fees'],.1)
    def test_missing_quote_not_fake_sale(self):
        p=self.opened(); self.md.available=False; self.exp.exit_one(p['id'])
        p=self.db.snapshot()['positions'][0]
        self.assertEqual(p['remaining_raw'],'4350000'); self.assertIsNone(p['mark_net'])
        self.assertIsNone(self.db.stats()[0]['equity'])
        self.assertEqual(self.db.allocation()[1],'UNPRICED_POSITION')
    def test_early_comparison_same_entry(self):
        self.opened(); self.md.multiple=1.3
        for p in self.db.snapshot()['positions']:self.exp.exit_one(p['id'])
        a,b=self.db.snapshot()['positions']
        self.assertEqual(a['initial_raw'],b['initial_raw'])
        self.assertFalse(a['half_done']); self.assertTrue(b['half_done'])
    def test_no_reentry(self):
        p=self.opened(); self.md.multiple=.8
        for x in self.db.snapshot()['positions']:self.exp.exit_one(x['id'])
        self.md.multiple=1
        self.assertEqual(self.exp.open_cohort('MINT','token',1,6),'ALREADY_TRADED')
    def test_changed_experiment_rejected(self):
        c=copy.deepcopy(CFG); c['risk_fraction']=.04
        with self.assertRaisesRegex(RuntimeError,'CONFIG_CHANGED'):Store(self.path,c)
    def test_export_safe_and_raw_exact(self):
        from openpyxl import load_workbook
        self.opened()
        book=load_workbook(io.BytesIO(export_excel(self.db)))
        sheet=book['Positions']; headers=[x.value for x in sheet[1]]
        self.assertEqual(sheet.cell(2,headers.index('symbol')+1).data_type,'s')
        self.assertEqual(sheet.cell(2,headers.index('initial_raw')+1).value,'4350000')
        self.assertIn('Events',book.sheetnames)
    def test_experiment_pause_persists_exits_continue(self):
        self.opened()
        with self.db.lock,self.db.cx:
            self.db.s['accounts']['PRINCIPAL_HALF']['cash']=100; self.db.save()
        self.db.sample_equity()
        self.assertTrue(self.db.snapshot()['paused'])
        p=self.db.snapshot()['positions'][0]; self.md.multiple=.8
        self.exp.exit_one(p['id'])
        self.assertEqual(self.db.snapshot()['positions'][0]['remaining_raw'],'0')
        self.assertEqual(self.db.allocation()[1],'EXPERIMENT_LOSS_PAUSE')

if __name__=='__main__':unittest.main(verbosity=2)
